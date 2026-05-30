"""导出 Excel / Word 月报、导入模板生成与解析"""
import io
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 批量导入字段映射：(中文列名, 数据库字段, 是否必填, 说明)
IMPORT_COLUMNS = [
    ('事故编号', 'accident_no', False, '留空自动生成，如 JT20250101001'),
    ('发生时间', 'occur_time', True, '格式：2025-01-15 14:30'),
    ('上报时间', 'report_time', False, '格式同上'),
    ('事故等级', 'level', False, '轻微事故/一般事故/重大事故/特大事故'),
    ('事故类型', 'type', False, '追尾/正面碰撞/侧面碰撞/刮擦/翻车/碾压/撞固定物/其他'),
    ('天气', 'weather', False, '晴/阴/雨/雪/雾/大风'),
    ('能见度', 'visibility', False, '良好/一般/较差/极差'),
    ('行政区划', 'district', False, ''),
    ('道路名称', 'road_name', False, ''),
    ('路段', 'road_section', False, ''),
    ('桩号', 'mileage', False, '如 K12+500'),
    ('经度', 'longitude', False, '数值，如 117.227'),
    ('纬度', 'latitude', False, '数值，如 31.820'),
    ('道路类型', 'road_type', False, '高速/国道/省道等'),
    ('道路线形', 'road_shape', False, '平直/弯道等'),
    ('路面状况', 'road_condition', False, '干燥/潮湿等'),
    ('死亡人数', 'death_count', False, '整数，默认 0'),
    ('受伤人数', 'injury_count', False, '整数，默认 0'),
    ('直接经济损失', 'economic_loss', False, '数值，单位元'),
    ('事故原因', 'cause', False, ''),
    ('处理结果', 'handle_result', False, '已结案/调解中等'),
    ('处理民警', 'handler', False, ''),
    ('备注', 'remarks', False, ''),
]


def export_accidents_excel(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = '事故列表'
    headers = ['事故编号', '发生时间', '等级', '类型', '行政区划', '道路名称', '路段',
               '天气', '路面状况', '死亡', '受伤', '直接损失(元)', '事故原因', '处理结果', '处理民警']
    ws.append(headers)
    head_fill = PatternFill('solid', fgColor='1f4e79')
    head_font = Font(bold=True, color='FFFFFF', size=11)
    thin = Side(border_style='thin', color='cccccc')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border
    widths = [18, 17, 10, 12, 12, 14, 10, 8, 10, 8, 8, 13, 14, 12, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    for r in rows:
        ws.append([
            r['accident_no'], r['occur_time'], r['level'], r['type'],
            r['district'], r['road_name'], r['road_section'],
            r['weather'], r['road_condition'],
            r['death_count'], r['injury_count'], r['economic_loss'],
            r['cause'], r['handle_result'], r['handler']
        ])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = border
            c.alignment = Alignment(vertical='center')
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_monthly_report_docx(year, month, summary, top_roads, type_dist, cause_dist, major_list):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f'连云港市东海县{year}年{month}月道路交通事故统计分析报告')
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1f, 0x4e, 0x79)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(f'报告生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M")}').font.size = Pt(10)

    doc.add_heading('一、事故总体概况', level=1)
    p = doc.add_paragraph()
    p.add_run(
        f'{year}年{month}月共发生道路交通事故 '
    )
    r = p.add_run(f'{summary["total"]} ')
    r.bold = True
    r.font.color.rgb = RGBColor(0xc0, 0x00, 0x00)
    p.add_run(
        f'起，造成死亡 {summary["deaths"]} 人，受伤 {summary["injuries"]} 人，'
        f'直接经济损失 {summary["loss"]:.2f} 元。'
    )
    if summary.get('total_yoy') is not None:
        yoy = summary['total_yoy']
        trend = '上升' if yoy > 0 else '下降' if yoy < 0 else '持平'
        p.add_run(f'与去年同期相比事故数{trend} {abs(yoy):.1f}%。')

    table = doc.add_table(rows=2, cols=4)
    table.style = 'Light Grid Accent 1'
    h = table.rows[0].cells
    h[0].text = '事故总数(起)'
    h[1].text = '死亡人数(人)'
    h[2].text = '受伤人数(人)'
    h[3].text = '直接损失(元)'
    v = table.rows[1].cells
    v[0].text = str(summary['total'])
    v[1].text = str(summary['deaths'])
    v[2].text = str(summary['injuries'])
    v[3].text = f'{summary["loss"]:.2f}'
    for c in h:
        for r in c.paragraphs:
            for run in r.runs:
                run.bold = True

    doc.add_heading('二、事故类型分布', level=1)
    if type_dist:
        t = doc.add_table(rows=1, cols=3)
        t.style = 'Light Grid Accent 1'
        hdr = t.rows[0].cells
        hdr[0].text = '事故类型'
        hdr[1].text = '数量(起)'
        hdr[2].text = '占比(%)'
        total = sum(x['count'] for x in type_dist) or 1
        for x in type_dist:
            row = t.add_row().cells
            row[0].text = x['name']
            row[1].text = str(x['count'])
            row[2].text = f'{x["count"] / total * 100:.1f}'
    else:
        doc.add_paragraph('（无数据）')

    doc.add_heading('三、事故主要原因分析', level=1)
    if cause_dist:
        t = doc.add_table(rows=1, cols=3)
        t.style = 'Light Grid Accent 1'
        hdr = t.rows[0].cells
        hdr[0].text = '事故原因'
        hdr[1].text = '数量(起)'
        hdr[2].text = '占比(%)'
        total = sum(x['count'] for x in cause_dist) or 1
        for x in cause_dist[:10]:
            row = t.add_row().cells
            row[0].text = x['name']
            row[1].text = str(x['count'])
            row[2].text = f'{x["count"] / total * 100:.1f}'
    else:
        doc.add_paragraph('（无数据）')

    doc.add_heading('四、事故多发路段 TOP10', level=1)
    if top_roads:
        t = doc.add_table(rows=1, cols=4)
        t.style = 'Light Grid Accent 1'
        hdr = t.rows[0].cells
        hdr[0].text = '排名'
        hdr[1].text = '道路名称'
        hdr[2].text = '事故数(起)'
        hdr[3].text = '死亡(人)'
        for i, x in enumerate(top_roads, 1):
            row = t.add_row().cells
            row[0].text = str(i)
            row[1].text = x['name']
            row[2].text = str(x['count'])
            row[3].text = str(x.get('deaths', 0))
    else:
        doc.add_paragraph('（无数据）')

    doc.add_heading('五、重大事故清单', level=1)
    if major_list:
        t = doc.add_table(rows=1, cols=5)
        t.style = 'Light Grid Accent 1'
        hdr = t.rows[0].cells
        hdr[0].text = '事故编号'
        hdr[1].text = '发生时间'
        hdr[2].text = '地点'
        hdr[3].text = '等级'
        hdr[4].text = '死/伤'
        for x in major_list:
            row = t.add_row().cells
            row[0].text = x['accident_no']
            row[1].text = x['occur_time']
            row[2].text = f"{x.get('road_name', '')} {x.get('road_section', '')}"
            row[3].text = x['level']
            row[4].text = f"{x['death_count']}/{x['injury_count']}"
    else:
        doc.add_paragraph('本月无重大及以上事故。')

    doc.add_heading('六、工作建议', level=1)
    suggestions = []
    if cause_dist:
        top_cause = cause_dist[0]['name']
        suggestions.append(f'本月事故主要原因为"{top_cause}"，建议加强针对性宣传教育与路面执法。')
    if top_roads:
        top_road = top_roads[0]['name']
        suggestions.append(f'"{top_road}"为事故高发路段，建议联合相关部门排查隐患、增设警示标识。')
    suggestions.append('持续推进酒驾醉驾、超速、违章变道等突出违法行为整治。')
    for s in suggestions:
        doc.add_paragraph(s, style='List Number')

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ============ 批量导入 ============

def build_import_template():
    """生成批量导入 Excel 模板（带表头、说明、示例行）"""
    wb = Workbook()
    ws = wb.active
    ws.title = '事故数据'
    headers = [c[0] for c in IMPORT_COLUMNS]
    ws.append(headers)

    head_fill = PatternFill('solid', fgColor='1f4e79')
    head_font = Font(bold=True, color='FFFFFF', size=11)
    req_fill = PatternFill('solid', fgColor='c0392b')
    thin = Side(border_style='thin', color='cccccc')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, c in enumerate(ws[1], 1):
        c.fill = req_fill if IMPORT_COLUMNS[i - 1][2] else head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = border
        desc = IMPORT_COLUMNS[i - 1][3]
        if desc:
            c.comment = Comment(desc, '系统')
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(IMPORT_COLUMNS[i - 1][0]) * 2 + 2)
    ws.row_dimensions[1].height = 26

    # 示例行
    ws.append([
        '', '2025-05-20 14:30', '2025-05-20 15:00',
        '一般事故', '追尾', '晴', '良好',
        '高新区', '人民路', '中段', 'K12+500',
        117.227, 31.820, '城市主干道', '平直', '干燥',
        0, 2, 8000.00,
        '未保持安全距离', '已结案', '王警官', '示例数据，请删除',
    ])

    # 说明 sheet
    ws2 = wb.create_sheet('填写说明')
    ws2.append(['列名', '是否必填', '说明'])
    for c in ws2[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal='center', vertical='center')
    for col in IMPORT_COLUMNS:
        ws2.append([col[0], '是' if col[2] else '否', col[3]])
    ws2.column_dimensions['A'].width = 16
    ws2.column_dimensions['B'].width = 10
    ws2.column_dimensions['C'].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _parse_dt(v):
    if v is None or v == '':
        return ''
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M')
    s = str(v).strip()
    for fmt in ('%Y-%m-%d %H:%M', '%Y-%m-%d %H:%M:%S', '%Y/%m/%d %H:%M', '%Y-%m-%d', '%Y/%m/%d'):
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d %H:%M')
        except ValueError:
            continue
    raise ValueError(f'无法识别的时间格式: {s}')


def parse_import_file(stream):
    """解析批量导入 Excel，返回 (records, errors)"""
    wb = load_workbook(stream, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], [{'row': 0, 'msg': '文件为空'}]
    header = [str(h).strip() if h else '' for h in rows[0]]
    name_to_field = {c[0]: c[1] for c in IMPORT_COLUMNS}
    required = {c[1] for c in IMPORT_COLUMNS if c[2]}
    # 列下标 -> 字段
    col_map = {}
    for i, name in enumerate(header):
        if name in name_to_field:
            col_map[i] = name_to_field[name]

    if not col_map:
        return [], [{'row': 1, 'msg': '未识别到任何已知列，请使用标准模板'}]

    records = []
    errors = []
    for ridx, row in enumerate(rows[1:], start=2):
        if not any(c not in (None, '') for c in row):
            continue
        rec = {'_row': ridx}
        try:
            for i, v in enumerate(row):
                if i not in col_map:
                    continue
                field = col_map[i]
                if v is None or v == '':
                    continue
                if field in ('death_count', 'injury_count'):
                    rec[field] = int(float(v))
                elif field in ('economic_loss', 'longitude', 'latitude'):
                    rec[field] = float(v)
                elif field in ('occur_time', 'report_time'):
                    rec[field] = _parse_dt(v)
                else:
                    rec[field] = str(v).strip()
            missing = [c[0] for c in IMPORT_COLUMNS if c[2] and not rec.get(c[1])]
            if missing:
                errors.append({'row': ridx, 'msg': f'缺少必填列: {", ".join(missing)}'})
                continue
            records.append(rec)
        except Exception as e:
            errors.append({'row': ridx, 'msg': f'数据格式错误: {e}'})
    return records, errors
